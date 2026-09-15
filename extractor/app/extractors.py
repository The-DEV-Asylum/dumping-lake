from pathlib import Path
from codecs import getincrementaldecoder
from typing import Iterator

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


READ_CHUNK = 1024 * 1024


def extract_plain(path: Path) -> Iterator[str]:
    decoder = getincrementaldecoder("utf-8")(errors="replace")
    with path.open("rb") as handle:
        while raw := handle.read(READ_CHUNK):
            text = decoder.decode(raw, final=False)
            if text:
                yield text
        tail = decoder.decode(b"", final=True)
        if tail:
            yield tail


def extract_pdf(path: Path) -> Iterator[str]:
    reader = PdfReader(str(path), strict=False)
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text:
            yield f"\n--- PAGE {number} ---\n{text}\n"


def extract_docx(path: Path) -> Iterator[str]:
    document = Document(str(path))
    for paragraph in document.paragraphs:
        if paragraph.text:
            yield paragraph.text + "\n"
    for table in document.tables:
        for row in table.rows:
            yield "\t".join(cell.text for cell in row.cells) + "\n"


def extract_xlsx(path: Path) -> Iterator[str]:
    workbook = load_workbook(str(path), read_only=True, data_only=False)
    try:
        for sheet in workbook.worksheets:
            yield f"\n--- SHEET {sheet.title} ---\n"
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                yield "\t".join(values) + "\n"
    finally:
        workbook.close()


def extract_html(path: Path) -> Iterator[str]:
    with path.open("rb") as handle:
        soup = BeautifulSoup(handle.read(), "html.parser")
    text = soup.get_text(" ", strip=True)
    if text:
        yield text + "\n"
